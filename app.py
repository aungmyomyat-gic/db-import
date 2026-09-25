import re
import os
import json
import shutil
import subprocess
import threading
import time
import unicodedata
import tempfile
import urllib.error
import urllib.request
import uuid
from datetime import date, datetime, time as datetime_time
from decimal import Decimal
from pathlib import Path

from flask import Flask, render_template, request, jsonify, session
import openpyxl
from openpyxl.utils import get_column_letter
import pyodbc

app = Flask(__name__)
app.secret_key = os.urandom(24)

# Persisted DB connection (port, username, password, database, host).
# Stored next to the app so it survives restarts; mount it as a volume to
# keep it across container rebuilds.
CONFIG_PATH = Path(os.environ.get("DB_CONFIG_PATH", "/app/data/db_config.json"))


def load_config() -> dict:
    if CONFIG_PATH.exists():
        try:
            return json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
        except Exception:
            return {}
    return {}


def save_config(cfg: dict) -> None:
    CONFIG_PATH.parent.mkdir(parents=True, exist_ok=True)
    CONFIG_PATH.write_text(json.dumps(cfg, indent=2, ensure_ascii=False), encoding="utf-8")

# A table name MUST be followed by a "(description)" — e.g. MZAIHP(倉庫在庫マスタ).
# This distinguishes table-name cells from header cells (DFRNKN, DFRGNO …),
# which are bare uppercase codes with no parentheses.
TABLE_NAME_RE = re.compile(r"^([A-Z][A-Z0-9]{3,9})\s*[\(（]")
HEADER_CODE_RE = re.compile(r"^[A-Z][A-Z0-9_]{1,30}$")

# Invisible characters that can hide at the start of a cell (zero-width, BOM).
_INVISIBLE_RE = re.compile(r"[​‌‍﻿ ]")


def normalize_cell(value: str) -> str:
    """
    Normalize a cell string so full-width letters/digits/parens typed in
    Japanese Excel (e.g. 'ＭＤＮＮＯＰ（…)') become plain ASCII ('MDNNOP(…)'),
    and strip invisible/zero-width characters that break matching.
    """
    s = unicodedata.normalize("NFKC", value)
    s = _INVISIBLE_RE.sub("", s)
    return s.strip()


# ── Excel helpers ─────────────────────────────────────────────────────────────

TARGET_SECTION = "実施前テストデータ"   # only import tables under this "■" section
MARKER_SCAN_WIDTH = 8

# ── Batch job trigger (derived from the uploaded file name) ────────────────────
BATCH_BASE_URL = os.environ.get("BATCH_BASE_URL", "http://localhost:8080")
BATCH_NAME_RE = re.compile(r"(day|month)[\s_-]*job[\s_-]*(\d+)", re.IGNORECASE)


def derive_batch_job(filename):
    """
    Map an uploaded file name to the batch endpoint it should trigger.
    'day_job_20.xlsx' / 'DayJob20.xlsx'   -> day-job20   / "Call DayJob20"
    'Month Job-1.xlsx' / 'monthjob1.xlsx' -> month-job1  / "Call Month Job-1"
    """
    if not filename:
        return None
    m = BATCH_NAME_RE.search(filename)
    if not m:
        return None
    kind, num = m.group(1).lower(), m.group(2)
    slug = f"{kind}-job{num}"
    label = f"Call DayJob{num}" if kind == "day" else f"Call Month Job-{num}"
    return {"slug": slug, "label": label}
IMPORT_BATCH_SIZE = int(os.environ.get("IMPORT_BATCH_SIZE", "5000"))
IMPORT_JOB_TTL_SECONDS = 3600
IMPORT_JOBS = {}
IMPORT_JOBS_LOCK = threading.Lock()


def is_one_line_marker_row(ws, row_num, start_col):
    values = []
    end_col = min(ws.max_column, start_col + MARKER_SCAN_WIDTH - 1)
    for col in range(start_col, end_col + 1):
        val = ws.cell(row=row_num, column=col).value
        if val is None:
            continue
        text = normalize_cell(str(val))
        if text:
            values.append(text)
            if len(values) > 1:
                return False
    return len(values) == 1 and not HEADER_CODE_RE.fullmatch(values[0])


def resolve_header_row(ws, meta):
    hrow = int(meta["header_row"])
    start_col = int(meta["start_col"])
    marker_skipped = is_one_line_marker_row(ws, hrow, start_col)
    return hrow + 1 if marker_skipped else hrow, marker_skipped


def scan_tables(ws):
    """
    Collect table blocks ONLY while inside the "■ 実施前テストデータ" section.
    Any other "■" section ("予想結果", "バッチ", …) is skipped entirely.
    If the sheet has no "■" markers at all, the whole sheet is scanned.
    """
    found, seen = [], set()
    in_target = None   # None = no marker seen yet (scan everything)
    seen_target = False
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        # A section marker is any cell beginning with "■"
        marker = next(
            (normalize_cell(value) for value in row
             if isinstance(value, str) and normalize_cell(value).startswith("■")),
            None,
        )
        if marker is not None:
            is_target = TARGET_SECTION in marker
            if seen_target and not is_target:
                break
            in_target = is_target
            seen_target = seen_target or is_target
            continue                 # skip the marker row itself

        if in_target is False:       # inside an excluded section → skip
            continue

        for col_idx, value in enumerate(row, start=1):
            if not value or not isinstance(value, str):
                continue
            m = TABLE_NAME_RE.match(normalize_cell(value))
            if not m:
                continue
            name = m.group(1)
            if name in seen:
                continue
            seen.add(name)
            found.append({
                "name": name,
                "name_row": row_idx,
                "start_col": col_idx,
                "header_row": row_idx + 1,
            })
    return found


def _row_value(row, col_num):
    idx = col_num - 1
    return row[idx] if idx < len(row) else None


def _is_one_line_marker_values(row, start_col):
    values = []
    end_col = start_col + MARKER_SCAN_WIDTH - 1
    for col in range(start_col, end_col + 1):
        val = _row_value(row, col)
        if val is None:
            continue
        text = normalize_cell(str(val))
        if text:
            values.append(text)
            if len(values) > 1:
                return False
    return len(values) == 1 and not HEADER_CODE_RE.fullmatch(values[0])


def _headers_from_values(row, start_col):
    headers, col = [], start_col
    while col <= len(row):
        val = _row_value(row, col)
        if val is None:
            break
        headers.append(normalize_cell(str(val)))
        col += 1
    return headers


def _preview_from_headers(meta, hrow, headers, marker_skipped):
    start_col = int(meta["start_col"])
    data_start_row = hrow + 1
    end_col = start_col + len(headers) - 1 if headers else start_col
    start_letter = get_column_letter(start_col)
    end_letter = get_column_letter(end_col)
    return {
        "name": meta["name"],
        "columns": len(headers),
        "records": 0,
        "name_row": int(meta["name_row"]),
        "header_row": hrow,
        "data_start_row": data_start_row,
        "data_end_row": None,
        "start_col": start_col,
        "end_col": end_col,
        "start_col_letter": start_letter,
        "end_col_letter": end_letter,
        "col_range": f"{start_letter}:{end_letter}",
        "row_range": "-",
        "cell_range": f"{start_letter}{hrow}:{end_letter}...",
        "marker_skipped": marker_skipped,
    }


def _count_active_preview_rows(active_tables, row, row_idx):
    still_active = []
    for table in active_tables:
        n_cols = table["columns"]
        start_col = table["start_col"]
        row_vals = [_row_value(row, start_col + i) for i in range(n_cols)]
        if all(v is None for v in row_vals):
            continue
        table["records"] += 1
        table["data_end_row"] = row_idx
        table["row_range"] = f"{table['data_start_row']}:{row_idx}"
        still_active.append(table)
    return still_active


def scan_table_previews(ws):
    """
    Stream the sheet once and build preview rows with real record counts.
    This avoids slow random access on read-only worksheets during preview.
    """
    tables, active, waiting = [], [], []
    seen = set()
    in_target = None
    seen_target = False

    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        active = _count_active_preview_rows(active, row, row_idx)

        next_waiting = []
        for item in waiting:
            meta = item["meta"]
            start_col = int(meta["start_col"])
            if row_idx == int(meta["header_row"]) and _is_one_line_marker_values(row, start_col):
                item["meta"] = {**meta, "header_row": row_idx + 1}
                item["marker_skipped"] = True
                next_waiting.append(item)
                continue

            headers = _headers_from_values(row, start_col)
            table = _preview_from_headers(meta, row_idx, headers, item["marker_skipped"])
            tables.append(table)
            if headers:
                active.append(table)
        waiting = next_waiting

        marker = next(
            (normalize_cell(value) for value in row
             if isinstance(value, str) and normalize_cell(value).startswith("■")),
            None,
        )
        if marker is not None:
            is_target = TARGET_SECTION in marker
            if seen_target and not is_target:
                break
            in_target = is_target
            seen_target = seen_target or is_target
            continue

        if in_target is False:
            continue

        for col_idx, value in enumerate(row, start=1):
            if not value or not isinstance(value, str):
                continue
            m = TABLE_NAME_RE.match(normalize_cell(value))
            if not m:
                continue
            name = m.group(1)
            if name in seen:
                continue
            seen.add(name)
            waiting.append({
                "meta": {
                    "name": name,
                    "name_row": row_idx,
                    "start_col": col_idx,
                    "header_row": row_idx + 1,
                },
                "marker_skipped": False,
            })

    return tables


def extract_headers(ws, meta):
    hrow, _ = resolve_header_row(ws, meta)
    start_col = int(meta["start_col"])

    headers, col = [], start_col
    while col <= ws.max_column:
        val = ws.cell(row=hrow, column=col).value
        if val is None:
            break
        # Normalize header codes (may be full-width) so they match DB columns.
        headers.append(normalize_cell(str(val)))
        col += 1
    end_col = start_col + len(headers) - 1 if headers else start_col
    return hrow, headers, end_col


def extract_table(ws, meta):
    hrow, headers, _ = extract_headers(ws, meta)
    start_col = int(meta["start_col"])
    if not headers:
        return [], []
    n_cols = len(headers)
    end_col = start_col + n_cols - 1
    rows = []
    # Sequential iter_rows() instead of per-cell ws.cell() lookups: on a
    # read_only workbook, random-access ws.cell() re-walks the sheet's XML
    # from the top on every call, making per-cell extraction O(rows^2).
    for row_vals in ws.iter_rows(
        min_row=hrow + 1, max_row=ws.max_row,
        min_col=start_col, max_col=end_col,
        values_only=True,
    ):
        if all(v is None for v in row_vals):
            break
        rows.append(row_vals)
    return headers, rows


# ── Routes ────────────────────────────────────────────────────────────────────

@app.route("/")
def index():
    return render_template("index.html")


@app.after_request
def disable_browser_cache(response):
    response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    return response


@app.route("/config", methods=["GET"])
def get_config():
    """Return the saved DB connection so the UI can pre-fill / skip setup."""
    cfg = load_config()
    return jsonify({"configured": bool(cfg), "config": cfg})


@app.route("/config", methods=["POST"])
def set_config():
    """Validate the connection, then persist it to JSON for later reuse."""
    data = request.json or {}
    host     = data.get("host", "").strip()
    port     = str(data.get("port", "")).strip()
    database = data.get("database", "").strip()
    username = data.get("username", "").strip()
    password = data.get("password", "")
    schema   = data.get("schema", "").strip() or "lvapdbf"

    if not all([host, port, database, username]):
        return jsonify({"error": "Host, port, database, and username are required."}), 400

    # Test the connection before saving so we never persist bad credentials.
    err = _test_connection(host, port, database, username, password)
    if err:
        return jsonify({"error": f"Connection failed: {err}"}), 400

    save_config({
        "host": host, "port": port, "database": database,
        "username": username, "password": password, "schema": schema,
    })
    return jsonify({"ok": True})


@app.route("/config", methods=["DELETE"])
def delete_config():
    """Remove the saved DB connection so the user can set up a new one."""
    if CONFIG_PATH.exists():
        CONFIG_PATH.unlink()
    return jsonify({"ok": True})


def _build_conn_str(host, port, database, username, password):
    drivers = [d for d in pyodbc.drivers() if "SQL Server" in d]
    if not drivers:
        return None, "No SQL Server ODBC driver found inside container."
    driver = sorted(drivers)[-1]
    conn_str = (
        f"DRIVER={{{driver}}};"
        f"SERVER={host},{port};"
        f"DATABASE={database};"
        f"UID={username};"
        f"PWD={password};"
        "TrustServerCertificate=yes;"
    )
    return conn_str, None


def _test_connection(host, port, database, username, password):
    conn_str, err = _build_conn_str(host, port, database, username, password)
    if err:
        return err
    try:
        pyodbc.connect(conn_str, timeout=10).close()
        return None
    except pyodbc.Error as e:
        return str(e)


@app.route("/upload", methods=["POST"])
def upload():
    """Save uploaded file, return sheet names."""
    file = request.files.get("file")
    if not file or not file.filename:
        return jsonify({"error": "No file selected."}), 400

    tmp = Path(tempfile.gettempdir()) / f"{uuid.uuid4()}.xlsx"
    file.save(tmp)

    try:
        wb = openpyxl.load_workbook(tmp, data_only=True, read_only=True)
        sheets = wb.sheetnames
        wb.close()
    except Exception as e:
        tmp.unlink(missing_ok=True)
        return jsonify({"error": f"Cannot open file: {e}"}), 400

    # Clean up previous temp file
    old = session.get("tmp_path")
    if old:
        Path(old).unlink(missing_ok=True)

    session["tmp_path"] = str(tmp)
    session["orig_filename"] = file.filename
    return jsonify({"sheets": sheets, "batch_job": derive_batch_job(file.filename)})


@app.route("/scan", methods=["POST"])
def scan():
    """Scan a sheet and return table preview."""
    payload = request.json or {}
    sheet_name = payload.get("sheet", "").strip()
    tmp_path = session.get("tmp_path")

    if not tmp_path or not Path(tmp_path).exists():
        return jsonify({"error": "Session expired — please re-upload the file."}), 400
    if not sheet_name:
        return jsonify({"error": "Sheet name is required."}), 400

    try:
        wb = openpyxl.load_workbook(tmp_path, data_only=True, read_only=True)
    except Exception as e:
        return jsonify({"error": f"Cannot open file: {e}"}), 400

    if sheet_name not in wb.sheetnames:
        return jsonify({"error": f"Sheet '{sheet_name}' not found."}), 400

    ws = wb[sheet_name]
    tables = scan_table_previews(ws)

    wb.close()
    session["sheet_name"] = sheet_name
    return jsonify({"tables": tables})


@app.route("/schemas", methods=["GET"])
def get_schemas():
    """List schemas that own user tables, so the UI can offer a target dropdown."""
    cfg = load_config()
    if not cfg:
        return jsonify({"error": "No DB connection saved."}), 400
    conn_str, err = _build_conn_str(
        cfg["host"], cfg["port"], cfg["database"], cfg["username"], cfg["password"]
    )
    if err:
        return jsonify({"error": err}), 500
    try:
        conn = pyodbc.connect(conn_str, timeout=10)
        cur = conn.cursor()
        cur.execute(
            "SELECT DISTINCT TABLE_SCHEMA FROM INFORMATION_SCHEMA.TABLES "
            "WHERE TABLE_TYPE = 'BASE TABLE' ORDER BY TABLE_SCHEMA"
        )
        schemas = [r[0] for r in cur.fetchall()]
        conn.close()
    except pyodbc.Error as e:
        return jsonify({"error": f"Could not list schemas: {e}"}), 400
    return jsonify({"schemas": schemas, "default": cfg.get("schema") or "lvapdbf"})


@app.route("/tables", methods=["GET"])
def get_tables():
    """List base tables in one schema for the maintenance UI."""
    cfg = load_config()
    if not cfg:
        return jsonify({"error": "No DB connection saved."}), 400
    schema = request.args.get("schema", "").strip()
    if not schema:
        return jsonify({"error": "Schema is required."}), 400

    conn_str, err = _build_conn_str(
        cfg["host"], cfg["port"], cfg["database"], cfg["username"], cfg["password"]
    )
    if err:
        return jsonify({"error": err}), 500
    try:
        conn = pyodbc.connect(conn_str, timeout=10)
        cur = conn.cursor()
        cur.execute(
            "SELECT TABLE_NAME FROM INFORMATION_SCHEMA.TABLES "
            "WHERE TABLE_TYPE = 'BASE TABLE' AND TABLE_SCHEMA = ? ORDER BY TABLE_NAME",
            schema,
        )
        tables = [row[0] for row in cur.fetchall()]
        conn.close()
    except pyodbc.Error as e:
        return jsonify({"error": f"Could not list tables: {e}"}), 400
    return jsonify({"schema": schema, "tables": tables})


def _table_metadata(cursor, schema, table_name):
    cursor.execute(
        "SELECT COLUMN_NAME, DATA_TYPE, IS_NULLABLE FROM INFORMATION_SCHEMA.COLUMNS "
        "WHERE TABLE_SCHEMA = ? AND TABLE_NAME = ? ORDER BY ORDINAL_POSITION",
        schema, table_name,
    )
    columns = [
        {"name": row[0], "type": row[1], "nullable": row[2] == "YES"}
        for row in cursor.fetchall()
    ]
    if not columns:
        return [], []
    cursor.execute(
        "SELECT k.COLUMN_NAME FROM INFORMATION_SCHEMA.TABLE_CONSTRAINTS t "
        "JOIN INFORMATION_SCHEMA.KEY_COLUMN_USAGE k "
        "ON t.CONSTRAINT_NAME = k.CONSTRAINT_NAME AND t.TABLE_SCHEMA = k.TABLE_SCHEMA "
        "WHERE t.CONSTRAINT_TYPE = 'PRIMARY KEY' AND t.TABLE_SCHEMA = ? AND t.TABLE_NAME = ? "
        "ORDER BY k.ORDINAL_POSITION",
        schema, table_name,
    )
    return columns, [row[0] for row in cursor.fetchall()]


def _json_table_value(value):
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (date, datetime, datetime_time)):
        return value.isoformat()
    if isinstance(value, bytes):
        return "0x" + value.hex()
    return str(value)


@app.route("/table-data", methods=["GET"])
def get_table_data():
    """Return a bounded table preview plus metadata for safe row editing."""
    cfg = load_config()
    if not cfg:
        return jsonify({"error": "No DB connection saved."}), 400
    schema = request.args.get("schema", "").strip()
    table_name = request.args.get("table", "").strip()
    limit = min(max(request.args.get("limit", 100, type=int), 1), 500)
    if not schema or not table_name:
        return jsonify({"error": "Schema and table are required."}), 400

    conn_str, err = _build_conn_str(
        cfg["host"], cfg["port"], cfg["database"], cfg["username"], cfg["password"]
    )
    if err:
        return jsonify({"error": err}), 500
    conn = None
    try:
        conn = pyodbc.connect(conn_str, timeout=10)
        cursor = conn.cursor()
        columns, primary_keys = _table_metadata(cursor, schema, table_name)
        if not columns:
            return jsonify({"error": f"Table '{schema}.{table_name}' was not found."}), 404
        safe_schema = schema.replace("]", "]]")
        safe_table = table_name.replace("]", "]]")
        order_sql = ""
        if primary_keys:
            order_sql = " ORDER BY " + ", ".join(f"[{key.replace(']', ']]')}]" for key in primary_keys)
        cursor.execute(f"SELECT TOP {limit} * FROM [{safe_schema}].[{safe_table}]{order_sql}")
        rows = [[_json_table_value(value) for value in row] for row in cursor.fetchall()]
        cursor.execute(f"SELECT COUNT_BIG(*) FROM [{safe_schema}].[{safe_table}]")
        total_rows = cursor.fetchone()[0]
        return jsonify({
            "schema": schema, "table": table_name, "columns": columns,
            "primary_keys": primary_keys,
            "identifier_columns": primary_keys or [column["name"] for column in columns],
            "rows": rows, "limit": limit, "total_rows": total_rows,
        })
    except pyodbc.Error as e:
        return jsonify({"error": f"Could not load table data: {e}"}), 400
    finally:
        if conn is not None:
            conn.close()


@app.route("/table-data", methods=["PATCH"])
def update_table_row():
    """Update one row identified by the table's primary key."""
    cfg = load_config()
    payload = request.json or {}
    schema = str(payload.get("schema", "")).strip()
    table_name = str(payload.get("table", "")).strip()
    keys = payload.get("keys") or {}
    values = payload.get("values") or {}
    if not cfg or not schema or not table_name or not isinstance(keys, dict) or not isinstance(values, dict):
        return jsonify({"error": "Invalid update request."}), 400
    if not values:
        return jsonify({"error": "No changed values were provided."}), 400

    conn_str, err = _build_conn_str(
        cfg["host"], cfg["port"], cfg["database"], cfg["username"], cfg["password"]
    )
    if err:
        return jsonify({"error": err}), 500
    conn = None
    try:
        conn = pyodbc.connect(conn_str, timeout=10)
        cursor = conn.cursor()
        columns, primary_keys = _table_metadata(cursor, schema, table_name)
        column_names = {column["name"] for column in columns}
        if not columns:
            return jsonify({"error": f"Table '{schema}.{table_name}' was not found."}), 404
        identifier_columns = primary_keys or [column["name"] for column in columns]
        if set(keys) != set(identifier_columns):
            return jsonify({"error": "Complete original row values are required to update this row."}), 400
        # Primary-key values may be changed; the original key values in `keys`
        # still identify the row in the WHERE clause.
        changed = [name for name in values if name in column_names]
        if len(changed) != len(values):
            return jsonify({"error": "One or more columns cannot be edited."}), 400

        quote = lambda name: "[" + name.replace("]", "]]") + "]"
        qualified = f"{quote(schema)}.{quote(table_name)}"
        set_sql = ", ".join(f"{quote(name)} = ?" for name in changed)
        where_parts = []
        where_params = []
        for name in identifier_columns:
            if keys[name] is None:
                where_parts.append(f"{quote(name)} IS NULL")
            else:
                where_parts.append(f"{quote(name)} = ?")
                where_params.append(keys[name])
        where_sql = " AND ".join(where_parts)
        params = [values[name] for name in changed] + where_params
        cursor.execute(f"UPDATE {qualified} SET {set_sql} WHERE {where_sql}", params)
        affected = cursor.rowcount
        if affected != 1:
            conn.rollback()
            return jsonify({"error": f"Expected one row, but matched {affected}. Data was not changed."}), 409
        conn.commit()
        return jsonify({"ok": True, "message": "Row updated successfully."})
    except pyodbc.Error as e:
        if conn is not None:
            conn.rollback()
        return jsonify({"error": f"Could not update row: {e}"}), 400
    finally:
        if conn is not None:
            conn.close()


# ── Version / update check ────────────────────────────────────────────────────
# version.json is baked into the image, so it is the version this copy runs.
# `update_url` (or the UPDATE_CHECK_URL env var) points at a published copy of
# the latest version.json; leaving it empty disables the check.
VERSION_PATH = Path(__file__).with_name("version.json")
UPDATE_CACHE_SECONDS = 30 * 60
_update_cache = {"at": 0.0, "data": None}


def _load_version_info() -> dict:
    try:
        return json.loads(VERSION_PATH.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {"version": "0.0.0"}


def _version_tuple(version: str) -> tuple:
    parts = [int(n) for n in re.findall(r"\d+", str(version))[:3]]
    return tuple(parts + [0] * (3 - len(parts)))


def _fetch_latest_version(url: str) -> dict:
    req = urllib.request.Request(url, headers={"Accept": "application/json", "Cache-Control": "no-cache"})
    with urllib.request.urlopen(req, timeout=5) as resp:
        return json.loads(resp.read().decode("utf-8"))


# ── Self-update (git pull inside the container) ───────────────────────────────
# docker-compose.yml mounts the project folder at /app, so pulling there
# updates the user's own clone; Flask debug mode reloads the code.
APP_DIR = Path(__file__).resolve().parent
DEFAULT_REPO_URL = "https://github.com/aungmyomyat-gic/db-import.git"
UPDATE_BRANCH = "main"
REBUILD_FILES = {"Dockerfile", "requirements.txt"}
_update_lock = threading.Lock()


def _git(*args, timeout=30):
    # Run git as the owner of the project folder so pulled files aren't left
    # owned by root on the host (Linux / WSL bind mounts).
    owner = (APP_DIR / ".git").stat()
    kwargs = {}
    if os.geteuid() == 0 and owner.st_uid != 0:
        kwargs = {"user": owner.st_uid, "group": owner.st_gid}
    env = {**os.environ, "HOME": "/tmp", "GIT_TERMINAL_PROMPT": "0"}
    return subprocess.run(
        ["git", "-c", f"safe.directory={APP_DIR}", "-C", str(APP_DIR), *args],
        capture_output=True, text=True, timeout=timeout, env=env, **kwargs,
    )


def _self_update_status():
    """Return (ok, reason) — whether the in-app Update button can work here."""
    if not shutil.which("git"):
        return False, "git is not installed in the container — rebuild once with docker compose up -d --build"
    if not (APP_DIR / ".git").exists():
        return False, "the app folder is not a git clone"
    branch = _git("rev-parse", "--abbrev-ref", "HEAD")
    if branch.returncode != 0:
        return False, branch.stderr.strip() or "could not read the git branch"
    if branch.stdout.strip() != UPDATE_BRANCH:
        return False, f"this copy is on branch '{branch.stdout.strip()}', updates only apply to '{UPDATE_BRANCH}'"
    return True, ""


@app.route("/update", methods=["POST"])
def run_self_update():
    # JSON-only so a plain cross-site form post can't trigger it
    if not request.is_json:
        return jsonify({"error": "Invalid update request."}), 400
    ok, reason = _self_update_status()
    if not ok:
        return jsonify({"error": f"Update is not available: {reason}."}), 400
    if not _update_lock.acquire(blocking=False):
        return jsonify({"error": "An update is already running."}), 409
    try:
        before = _git("rev-parse", "HEAD").stdout.strip()
        repo_url = _load_version_info().get("repo_url") or DEFAULT_REPO_URL
        pull = _git("pull", "--ff-only", repo_url, UPDATE_BRANCH, timeout=120)
        output = (pull.stdout + pull.stderr).strip()
        if pull.returncode != 0:
            return jsonify({"error": "git pull failed.", "output": output}), 400
        after = _git("rev-parse", "HEAD").stdout.strip()
        changed = []
        if before != after:
            diff = _git("diff", "--name-only", before, after)
            changed = [line for line in diff.stdout.splitlines() if line]
        return jsonify({
            "ok": True, "updated": before != after, "output": output,
            "version": _load_version_info().get("version"),
            "needs_rebuild": any(Path(f).name in REBUILD_FILES for f in changed),
        })
    except (OSError, subprocess.SubprocessError) as e:
        return jsonify({"error": f"Update failed: {e}"}), 500
    finally:
        _update_lock.release()


@app.route("/version", methods=["GET"])
def get_version():
    local = _load_version_info()
    can_update, update_reason = _self_update_status()
    result = {
        "current": local.get("version", "0.0.0"), "date": local.get("date"), "notes": local.get("notes", []),
        "can_self_update": can_update, "self_update_reason": update_reason,
    }
    url = os.environ.get("UPDATE_CHECK_URL") or local.get("update_url") or ""
    if not url:
        return jsonify({**result, "check": "disabled"})

    now = time.time()
    force = request.args.get("force") == "1"
    if force or _update_cache["data"] is None or now - _update_cache["at"] > UPDATE_CACHE_SECONDS:
        try:
            _update_cache["data"] = _fetch_latest_version(url)
            _update_cache["at"] = now
        except (OSError, ValueError, urllib.error.URLError) as e:
            return jsonify({**result, "check": "failed", "error": str(e)})
    latest = _update_cache["data"] or {}
    latest_version = str(latest.get("version", "0.0.0"))
    return jsonify({
        **result, "check": "ok", "latest": latest_version,
        "latest_date": latest.get("date"), "latest_notes": latest.get("notes", []),
        "update_available": _version_tuple(latest_version) > _version_tuple(result["current"]),
    })


@app.route("/table-data", methods=["POST"])
def insert_table_rows():
    """Insert pasted rows into a table in a single all-or-nothing transaction."""
    cfg = load_config()
    payload = request.json or {}
    schema = str(payload.get("schema", "")).strip()
    table_name = str(payload.get("table", "")).strip()
    names = payload.get("columns") or []
    rows = payload.get("rows") or []
    if not cfg or not schema or not table_name or not isinstance(names, list) or not isinstance(rows, list):
        return jsonify({"error": "Invalid insert request."}), 400
    if not names or not rows:
        return jsonify({"error": "No rows to insert."}), 400
    if len(rows) > 5000:
        return jsonify({"error": "Paste at most 5000 rows at a time."}), 400
    if len(set(names)) != len(names) or any(not isinstance(r, list) or len(r) != len(names) for r in rows):
        return jsonify({"error": "Every row must have one value per column."}), 400

    conn_str, err = _build_conn_str(
        cfg["host"], cfg["port"], cfg["database"], cfg["username"], cfg["password"]
    )
    if err:
        return jsonify({"error": err}), 500
    conn = None
    try:
        conn = pyodbc.connect(conn_str, timeout=10)
        cursor = conn.cursor()
        columns, _ = _table_metadata(cursor, schema, table_name)
        if not columns:
            return jsonify({"error": f"Table '{schema}.{table_name}' was not found."}), 404
        by_name = {column["name"]: column for column in columns}
        unknown = [name for name in names if name not in by_name]
        if unknown:
            return jsonify({"error": f"Unknown column(s): {', '.join(unknown)}"}), 400

        quote = lambda name: "[" + name.replace("]", "]]") + "]"
        qualified = f"{quote(schema)}.{quote(table_name)}"
        cursor.execute(
            "SELECT name FROM sys.columns WHERE object_id = OBJECT_ID(?) AND is_identity = 1", qualified
        )
        identity = {row[0] for row in cursor.fetchall()}
        prefix = f"SET IDENTITY_INSERT {qualified} ON; " if identity & set(names) else ""
        suffix = f"; SET IDENTITY_INSERT {qualified} OFF" if prefix else ""
        insert_sql = (
            f"{prefix}INSERT INTO {qualified} ({', '.join(quote(n) for n in names)}) "
            f"VALUES ({', '.join('?' for _ in names)}){suffix}"
        )

        for row_number, row in enumerate(rows, start=1):
            values = []
            for name, value in zip(names, row):
                # Empty Excel cells become NULL, or the type default for NOT NULL columns
                if value is None or (isinstance(value, str) and value == ""):
                    value = None if by_name[name]["nullable"] else _type_default(by_name[name]["type"])
                values.append(value)
            try:
                cursor.execute(insert_sql, values)
            except pyodbc.Error as e:
                conn.rollback()
                return jsonify({"error": f"Row {row_number} failed, nothing was inserted: {e}"}), 400
        conn.commit()
        return jsonify({"ok": True, "inserted": len(rows), "message": f"{len(rows)} row(s) inserted."})
    except pyodbc.Error as e:
        if conn is not None:
            conn.rollback()
        return jsonify({"error": f"Could not insert rows: {e}"}), 400
    finally:
        if conn is not None:
            conn.close()


@app.route("/import", methods=["POST"])
def do_import():
    """Start a background import job and return its initial status."""
    cfg = load_config()
    if not cfg:
        return jsonify({"error": "No DB connection saved. Please set it up first."}), 400

    tmp_path   = session.get("tmp_path")
    sheet_name = session.get("sheet_name")
    payload = request.json or {}

    # Only import the tables the user ticked (None = all, for backward compat).
    selected = payload.get("tables")
    selected_names = None
    if selected is not None:
        selected_names = [str(name) for name in selected if str(name).strip()]
    if selected_names is not None and not selected_names:
        return jsonify({"error": "No tables selected."}), 400

    if not tmp_path or not Path(tmp_path).exists():
        return jsonify({"error": "Session expired — please re-upload the file."}), 400

    conn_str, err = _build_conn_str(
        cfg["host"], cfg["port"], cfg["database"], cfg["username"], cfg["password"]
    )
    if err:
        return jsonify({"error": err}), 500

    # Schema chosen in the UI dropdown wins; fall back to saved config, then default.
    preferred_schema = payload.get("schema", "").strip() \
        or cfg.get("schema") or "lvapdbf"

    _cleanup_import_jobs()
    job_id = uuid.uuid4().hex
    job = _new_import_job(job_id, selected_names or [])
    with IMPORT_JOBS_LOCK:
        IMPORT_JOBS[job_id] = job
        snapshot = _job_snapshot(job)

    worker = threading.Thread(
        target=_run_import_job,
        args=(job_id, cfg, str(tmp_path), sheet_name, selected_names, preferred_schema),
        daemon=True,
    )
    worker.start()
    return jsonify(snapshot), 202


@app.route("/import/status/<job_id>", methods=["GET"])
def import_status(job_id):
    _cleanup_import_jobs()
    with IMPORT_JOBS_LOCK:
        job = IMPORT_JOBS.get(job_id)
        if not job:
            return jsonify({"error": "Import job not found."}), 404
        return jsonify(_job_snapshot(job))


@app.route("/truncate", methods=["POST"])
def truncate_table():
    """Remove all rows from one validated table in the selected schema."""
    cfg = load_config()
    if not cfg:
        return jsonify({"error": "No DB connection saved. Please set it up first."}), 400

    payload = request.json or {}
    table_name = str(payload.get("table", "")).strip()
    schema = str(payload.get("schema", "")).strip() or cfg.get("schema") or "lvapdbf"
    if not table_name or not schema:
        return jsonify({"error": "Table and schema are required."}), 400

    conn_str, err = _build_conn_str(
        cfg["host"], cfg["port"], cfg["database"], cfg["username"], cfg["password"]
    )
    if err:
        return jsonify({"error": err}), 500

    conn = None
    try:
        conn = pyodbc.connect(conn_str, timeout=10)
        cursor = conn.cursor()
        cursor.execute(
            "SELECT 1 FROM INFORMATION_SCHEMA.TABLES "
            "WHERE TABLE_TYPE = 'BASE TABLE' AND TABLE_SCHEMA = ? AND TABLE_NAME = ?",
            schema, table_name,
        )
        if cursor.fetchone() is None:
            return jsonify({"error": f"Table '{schema}.{table_name}' was not found."}), 404

        safe_schema = schema.replace("]", "]]")
        safe_table = table_name.replace("]", "]]")
        cursor.execute(f"TRUNCATE TABLE [{safe_schema}].[{safe_table}]")
        conn.commit()
        return jsonify({"ok": True, "message": f"{schema}.{table_name} truncated successfully."})
    except pyodbc.Error as e:
        if conn is not None:
            conn.rollback()
        return jsonify({"error": f"Could not truncate '{schema}.{table_name}': {e}"}), 400
    finally:
        if conn is not None:
            conn.close()


@app.route("/call-batch", methods=["POST"])
def call_batch():
    """Trigger the batch endpoint matching the uploaded file's name."""
    job = derive_batch_job(session.get("orig_filename"))
    if not job:
        return jsonify({"error": "Could not determine a batch job from the file name."}), 400

    url = f"{BATCH_BASE_URL}/batch/{job['slug']}"
    try:
        with urllib.request.urlopen(url, timeout=15) as resp:
            return jsonify({"ok": True, "slug": job["slug"], "status": resp.status})
    except urllib.error.HTTPError as e:
        return jsonify({"error": f"{url} responded {e.code}: {e.reason}"}), 502
    except urllib.error.URLError as e:
        return jsonify({"error": f"Could not reach {url}: {e.reason}"}), 502


def _new_import_job(job_id, table_names):
    now = time.time()
    return {
        "id": job_id,
        "state": "queued",
        "total": len(table_names),
        "completed": 0,
        "results": [
            {"name": name, "status": "queued", "message": "Waiting", "errors": [],
             "progress": {"done": 0, "total": 0}}
            for name in table_names
        ],
        "error": None,
        "created_at": now,
        "updated_at": now,
    }


def _job_snapshot(job):
    return {
        "job_id": job["id"],
        "state": job["state"],
        "total": job["total"],
        "completed": job["completed"],
        "results": [dict(result) for result in job["results"]],
        "error": job.get("error"),
    }


def _cleanup_import_jobs():
    cutoff = time.time() - IMPORT_JOB_TTL_SECONDS
    with IMPORT_JOBS_LOCK:
        for job_id, job in list(IMPORT_JOBS.items()):
            if job["state"] in {"done", "failed"} and job["updated_at"] < cutoff:
                del IMPORT_JOBS[job_id]


def _replace_job_queue(job_id, table_names):
    now = time.time()
    with IMPORT_JOBS_LOCK:
        job = IMPORT_JOBS.get(job_id)
        if not job:
            return
        job["total"] = len(table_names)
        job["completed"] = 0
        job["results"] = [
            {"name": name, "status": "queued", "message": "Waiting", "errors": [],
             "progress": {"done": 0, "total": 0}}
            for name in table_names
        ]
        job["updated_at"] = now


def _set_job_state(job_id, state, error=None):
    with IMPORT_JOBS_LOCK:
        job = IMPORT_JOBS.get(job_id)
        if not job:
            return
        job["state"] = state
        job["error"] = error
        job["updated_at"] = time.time()


def _mark_job_table(job_id, table_name, status, message, errors=None):
    final_statuses = {"ok", "warn", "skip", "failed"}
    with IMPORT_JOBS_LOCK:
        job = IMPORT_JOBS.get(job_id)
        if not job:
            return
        existing = next((r for r in job["results"] if r["name"] == table_name), None)
        if existing is None:
            existing = {"name": table_name, "status": "queued", "message": "Waiting", "errors": [],
                        "progress": {"done": 0, "total": 0}}
            job["results"].append(existing)
            job["total"] = max(job["total"], len(job["results"]))

        was_final = existing["status"] in final_statuses
        existing.update({
            "status": status,
            "message": message,
            "errors": errors or [],
        })
        if status in final_statuses and not was_final:
            job["completed"] += 1
        job["updated_at"] = time.time()


def _mark_job_progress(job_id, table_name, done, total):
    with IMPORT_JOBS_LOCK:
        job = IMPORT_JOBS.get(job_id)
        if not job:
            return
        existing = next((r for r in job["results"] if r["name"] == table_name), None)
        if existing is None:
            return
        existing["progress"] = {"done": done, "total": total}
        job["updated_at"] = time.time()


def _finish_import_job(job_id):
    with IMPORT_JOBS_LOCK:
        job = IMPORT_JOBS.get(job_id)
        if not job:
            return
        if job["state"] != "failed":
            job["state"] = "done"
        job["updated_at"] = time.time()


def _fail_import_job(job_id, message):
    with IMPORT_JOBS_LOCK:
        job = IMPORT_JOBS.get(job_id)
        if not job:
            return
        final_statuses = {"ok", "warn", "skip", "failed"}
        for result in job["results"]:
            if result["status"] not in final_statuses:
                result["status"] = "failed"
                result["message"] = message
                result["errors"] = []
                job["completed"] += 1
        job["state"] = "failed"
        job["error"] = message
        job["updated_at"] = time.time()


def _run_import_job(job_id, cfg, tmp_path, sheet_name, selected_names, preferred_schema):
    conn = None
    wb = None
    try:
        _set_job_state(job_id, "running")
        conn_str, err = _build_conn_str(
            cfg["host"], cfg["port"], cfg["database"], cfg["username"], cfg["password"]
        )
        if err:
            _fail_import_job(job_id, err)
            return
        conn = pyodbc.connect(conn_str, timeout=10)

        wb = openpyxl.load_workbook(tmp_path, data_only=True, read_only=True)
        if sheet_name not in wb.sheetnames:
            _fail_import_job(job_id, f"Sheet '{sheet_name}' not found.")
            return

        ws = wb[sheet_name]
        metas = scan_tables(ws)
        if selected_names is None:
            selected_names = [meta["name"] for meta in metas]
            _replace_job_queue(job_id, selected_names)

        selected_set = set(selected_names)
        pending = set(selected_names)
        for meta in metas:
            table_name = meta["name"]
            if table_name not in selected_set:
                continue
            pending.discard(table_name)
            _mark_job_table(job_id, table_name, "running", "Importing")
            try:
                headers, rows = extract_table(ws, meta)
                if not headers:
                    result = {"name": table_name, "status": "skip", "message": "Header row empty", "errors": []}
                else:
                    progress_cb = lambda done, total, tn=table_name: _mark_job_progress(job_id, tn, done, total)
                    result = _import_one(conn, table_name, headers, rows, preferred_schema, progress_cb)
            except Exception as e:
                try:
                    conn.rollback()
                except Exception:
                    pass
                result = {"name": table_name, "status": "failed", "message": f"Failed: {e}", "errors": []}
            _mark_job_table(
                job_id,
                table_name,
                result.get("status", "failed"),
                result.get("message", ""),
                result.get("errors", []),
            )

        for table_name in pending:
            _mark_job_table(job_id, table_name, "skip", "Table not found in sheet", [])
        _finish_import_job(job_id)
    except Exception as e:
        _fail_import_job(job_id, f"Import failed: {e}")
    finally:
        if conn is not None:
            try:
                conn.close()
            except Exception:
                pass
        if wb is not None:
            try:
                wb.close()
            except Exception:
                pass


def _values_for_insert(row, plan, db_cols):
    values = []
    for col_name, src in plan:
        if src[0] == "sheet":
            v = row[src[1]]
            if v is None and not db_cols[col_name.upper()]["nullable"]:
                v = _type_default(db_cols[col_name.upper()]["type"])
            values.append(v)
        else:  # "fill"
            values.append(src[1])
    return tuple(values)


def _insert_batch(cursor, insert_sql, batch, error_samples):
    cursor.execute("SAVE TRANSACTION import_batch")
    try:
        cursor.fast_executemany = True
        cursor.executemany(insert_sql, batch)
        return len(batch), 0
    except pyodbc.Error:
        cursor.fast_executemany = False
        cursor.execute("ROLLBACK TRANSACTION import_batch")

    inserted = errors = 0
    for values in batch:
        try:
            cursor.execute(insert_sql, values)
            inserted += 1
        except pyodbc.Error as e:
            errors += 1
            if len(error_samples) < 2:
                error_samples.append(str(e))
    return inserted, errors


def _import_one(conn, table_name, headers, rows, preferred_schema="lvapdbf", progress_cb=None):
    cursor = conn.cursor()

    # Resolve the actual schema (e.g. 'lvapdbf'), not the login default ('dbo').
    cursor.execute(
        "SELECT TABLE_SCHEMA FROM INFORMATION_SCHEMA.TABLES WHERE TABLE_NAME = ?",
        table_name,
    )
    schemas = [r[0] for r in cursor.fetchall()]
    if not schemas:
        return {"name": table_name, "status": "skip", "message": "Table not found in DB"}
    # Import into the chosen schema; if the table doesn't exist there, skip it.
    if preferred_schema not in schemas:
        return {"name": table_name, "status": "skip",
                "message": f"Not in schema '{preferred_schema}' (found in: {', '.join(schemas)})"}
    schema = preferred_schema
    qualified = f"[{schema}].[{table_name}]"

    # Full column metadata: type + nullability + whether it has a DB default.
    cursor.execute(
        "SELECT COLUMN_NAME, DATA_TYPE, IS_NULLABLE, COLUMN_DEFAULT "
        "FROM INFORMATION_SCHEMA.COLUMNS "
        "WHERE TABLE_NAME = ? AND TABLE_SCHEMA = ? ORDER BY ORDINAL_POSITION",
        table_name, schema,
    )
    db_cols = {}
    for name, dtype, nullable, default in cursor.fetchall():
        db_cols[name.upper()] = {
            "name": name,
            "type": (dtype or "").lower(),
            "nullable": nullable == "YES",
            "has_default": default is not None,
        }

    # Which sheet header maps to which DB column?
    sheet_idx = {h.upper(): i for i, h in enumerate(headers) if h.upper() in db_cols}
    if not sheet_idx:
        return {"name": table_name, "status": "skip", "message": "No matching columns"}

    # Build the column list to INSERT:
    #  • sheet columns      → take the cell value (NULL→default if NOT NULL)
    #  • missing NOT NULL    → supply a type default (unless the DB has its own default)
    #  • everything else     → omit (let DB default / NULL apply)
    plan = []   # (db_col_name, source)  source = ("sheet", idx) | ("fill", value)
    filled_cols = []
    for up, info in db_cols.items():
        if up in sheet_idx:
            plan.append((info["name"], ("sheet", sheet_idx[up])))
        elif not info["nullable"] and not info["has_default"]:
            plan.append((info["name"], ("fill", _type_default(info["type"]))))
            filled_cols.append(info["name"])

    # TRUNCATE is minimally logged and skips per-row delete cost; fall back to
    # DELETE only if it's blocked (e.g. a foreign key references this table).
    try:
        cursor.execute(f"TRUNCATE TABLE {qualified}")
    except pyodbc.Error:
        conn.rollback()
        cursor.execute(f"DELETE FROM {qualified}")

    cols_sql     = ", ".join(f"[{c}]" for c, _ in plan)
    placeholders = ", ".join("?" for _ in plan)
    insert_sql   = f"INSERT INTO {qualified} ({cols_sql}) VALUES ({placeholders})"

    total_rows = len(rows)
    inserted = errors = 0
    error_samples = []
    batch = []
    if progress_cb:
        progress_cb(0, total_rows)
    for row in rows:
        batch.append(_values_for_insert(row, plan, db_cols))
        if len(batch) >= IMPORT_BATCH_SIZE:
            batch_inserted, batch_errors = _insert_batch(cursor, insert_sql, batch, error_samples)
            inserted += batch_inserted
            errors += batch_errors
            batch = []
            if progress_cb:
                progress_cb(inserted + errors, total_rows)
    if batch:
        batch_inserted, batch_errors = _insert_batch(cursor, insert_sql, batch, error_samples)
        inserted += batch_inserted
        errors += batch_errors

    conn.commit()
    if progress_cb:
        progress_cb(inserted + errors, total_rows)
    unmapped = len(headers) - len(sheet_idx)
    msg = f"{inserted} rows inserted"
    if errors:
        msg += f", {errors} errors"
    if filled_cols:
        msg += f" ({len(filled_cols)} required cols auto-filled: {', '.join(filled_cols)})"
    if unmapped:
        msg += f" ({unmapped} sheet cols skipped — not in table)"

    return {
        "name": table_name,
        "status": "ok" if not errors else "warn",
        "message": msg,
        "inserted": inserted,
        "errors": error_samples,
    }


def _type_default(dtype: str):
    """A safe non-NULL default for a SQL Server column type."""
    numeric = {
        "int", "bigint", "smallint", "tinyint", "bit",
        "decimal", "numeric", "money", "smallmoney", "float", "real",
    }
    dates = {"date", "datetime", "datetime2", "smalldatetime", "datetimeoffset"}
    if dtype in numeric:
        return 0
    if dtype == "time":
        return "00:00:00"
    if dtype in dates:
        return "1900-01-01"
    return ""   # char / varchar / nchar / nvarchar / text → empty string


if __name__ == "__main__":
    debug = os.environ.get("FLASK_DEBUG", "").lower() in {"1", "true", "yes"}
    app.run(host="0.0.0.0", port=5000, debug=debug)
